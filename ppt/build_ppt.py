# -*- coding: utf-8 -*-
"""Builds Bus_Bunching_Control_System.pptx.

Content basis: Copy-of-Bus-Bunching-Control-System.pdf (10 slides).
Design/format basis: Copy-of-Energy-Constrained-Dynamic-Travelling-Salesman-
Problem-for-Electric-Vehicle-Fleets.pdf (white bg, black geometric-sans
headings, green/orange accents, numbered "01" + underline card pattern,
mint formula boxes, dark-navy vs. light two-column comparison, horizontal
timeline diagrams).
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE

HERE = os.path.dirname(os.path.abspath(__file__))
EXTRACTED = os.path.join(HERE, "assets")
OUT = os.path.join(HERE, "Bus_Bunching_Control_System.pptx")

# ---------- palette (sampled from the EV-fleet design-reference deck) ------
INK    = RGBColor(0x16, 0x17, 0x1A)   # headings
BODY   = RGBColor(0x5B, 0x5F, 0x66)   # body / gray text
GREEN  = RGBColor(0x1E, 0x9E, 0x6B)   # accent green
ORANGE = RGBColor(0xE2, 0x79, 0x2D)   # accent orange
NAVY   = RGBColor(0x1E, 0x20, 0x24)   # dark comparison box
CARD   = RGBColor(0xF5, 0xF4, 0xF1)   # light warm-gray card bg
MINT   = RGBColor(0xDE, 0xF3, 0xE7)   # formula box bg
BORDER = RGBColor(0xDA, 0xDA, 0xD8)   # thin borders
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GRAYLT = RGBColor(0x9A, 0x9D, 0xA1)   # faint labels (page numbers, "01")

FONT_HEAD = "Segoe UI Semibold"
FONT_BODY = "Segoe UI"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]

_page_no = [0]


def new_slide(bg=WHITE):
    slide = prs.slides.add_slide(BLANK)
    r = slide.background.fill
    r.solid()
    r.fore_color.rgb = bg
    _page_no[0] += 1
    if _page_no[0] > 1:
        add_text(slide, SLIDE_W - Inches(1.0), SLIDE_H - Inches(0.5),
                  Inches(0.7), Inches(0.35), str(_page_no[0]), 10,
                  color=GRAYLT, align=PP_ALIGN.RIGHT)
    return slide


def _set_run(run, text, size, bold=False, color=INK, italic=False,
             name=FONT_BODY):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = name
    run.font.color.rgb = color


def add_text(slide, left, top, width, height, text, size, bold=False,
             color=INK, align=PP_ALIGN.LEFT, italic=False, name=FONT_BODY,
             line_spacing=1.0, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        _set_run(run, line, size, bold, color, italic, name)
    return tb


def add_bullets(slide, left, top, width, height, items, size=13,
                 color=BODY, bullet_color=None, gap=8, bold_lead=None,
                 name=FONT_BODY, bullet_char="•  "):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    bullet_color = bullet_color or color
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.12
        r1 = p.add_run()
        _set_run(r1, bullet_char, size, False, bullet_color, name=name)
        r2 = p.add_run()
        _set_run(r2, item, size, False, color, name=name)
    return tb


def rounded_rect(slide, left, top, w, h, fill=None, line=None,
                  line_w=Pt(1), radius=0.06, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w
    shp.shadow.inherit = False
    return shp


def rect(slide, left, top, w, h, fill=None, line=None, line_w=Pt(1)):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w
    shp.shadow.inherit = False
    return shp


def oval(slide, left, top, w, h, fill=None, line=None, line_w=Pt(1.25)):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w
    shp.shadow.inherit = False
    return shp


def hline(slide, left, top, w, color=BORDER, weight=Pt(1.5)):
    shp = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, left, top,
                                      left + w, top)
    shp.line.color.rgb = color
    shp.line.width = weight
    shp.shadow.inherit = False
    return shp


def icon_badge(slide, cx, cy, d, symbol, fill=CARD, sym_color=INK,
                line=None):
    o = oval(slide, cx - d // 2, cy - d // 2, d, d, fill=fill, line=line,
             line_w=Pt(1.25))
    tb = slide.shapes.add_textbox(cx - d // 2, cy - d // 2, d, d)
    tf = tb.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _set_run(r, symbol, int(d.pt * 0.34), color=sym_color,
             name="Segoe UI Symbol")
    return o


def add_title(slide, text, size=30, top=Inches(0.5), left=Inches(0.75),
              width=Inches(11.8)):
    return add_text(slide, left, top, width, Inches(0.9), text, size,
                     bold=True, color=INK, name=FONT_HEAD)


MX = Inches(0.75)  # left/right content margin


# ---------------------------------------------------------------------------
# SLIDE 1 — Title
# ---------------------------------------------------------------------------
def slide_title():
    s = new_slide()
    panel_w = Inches(4.7)
    rect(s, 0, 0, panel_w, SLIDE_H, fill=CARD)

    # simple flat "route" motif: dotted curved path between stop markers
    cx = Inches(0.6)
    pts_y = [Inches(5.7), Inches(4.5), Inches(3.1), Inches(1.7)]
    pts_x = [Inches(1.0), Inches(2.6), Inches(1.6), Inches(3.4)]
    for i in range(len(pts_x) - 1):
        connector = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, pts_x[i],
                                            pts_y[i], pts_x[i + 1], pts_y[i + 1])
        connector.line.color.rgb = GREEN
        connector.line.width = Pt(2.25)
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        connector.shadow.inherit = False
    stop_d = Inches(0.22)
    for i, (x, y) in enumerate(zip(pts_x, pts_y)):
        col = NAVY if i in (0, len(pts_x) - 1) else GREEN
        oval(s, x - stop_d // 2, y - stop_d // 2, stop_d, stop_d,
             fill=WHITE, line=col, line_w=Pt(2.5))
    # a small flat "bus" icon (body + two wheels), drawn as plain shapes
    bus_w, bus_h = Inches(0.6), Inches(0.32)
    bx, by = pts_x[-1] - bus_w // 2, pts_y[-1] - Inches(0.75)
    rounded_rect(s, bx, by, bus_w, bus_h, fill=NAVY, radius=0.3)
    win_w, win_h = Inches(0.1), Inches(0.14)
    for wi in range(3):
        wx = bx + Inches(0.09) + wi * Inches(0.15)
        oval(s, wx, by + Inches(0.06), win_w, win_h, fill=WHITE)
    wheel_d = Inches(0.11)
    oval(s, bx + Inches(0.08), by + bus_h - wheel_d // 2, wheel_d, wheel_d,
         fill=INK)
    oval(s, bx + bus_w - Inches(0.19), by + bus_h - wheel_d // 2, wheel_d,
         wheel_d, fill=INK)

    tx = panel_w + Inches(0.7)
    tw = SLIDE_W - tx - Inches(0.7)
    add_text(s, tx, Inches(1.55), tw, Inches(1.9),
             "Bus Bunching Control System", 40, bold=True, color=INK,
             name=FONT_HEAD, line_spacing=1.03)
    add_text(s, tx, Inches(2.95), tw, Inches(0.6),
             "Real-Time Bus Tracking, Route Diversion, and Anti-Bunching Dashboard",
             15, bold=True, color=BODY, line_spacing=1.15)
    add_text(s, tx, Inches(3.55), tw, Inches(0.4),
             "Under the Guidance of Dr. Anand S. Krishna", 13, color=BODY,
             italic=True)
    add_bullets(s, tx, Inches(4.15), tw, Inches(1.8),
                ["Deepti V Ligammanavar", "Gowda H L K", "Khalandar S",
                 "Roushni Zynab"], size=13, color=INK, bullet_color=GREEN,
                gap=6)


# ---------------------------------------------------------------------------
# SLIDE 2 — Problem Statement
# ---------------------------------------------------------------------------
def slide_problem():
    s = new_slide()
    add_title(s, "Problem Statement")
    box_top = Inches(2.0)
    box = rounded_rect(s, MX, box_top, SLIDE_W - 2 * MX, Inches(1.9),
                        fill=WHITE, line=BORDER, line_w=Pt(1), radius=0.04)
    bar = rect(s, MX, box_top, Inches(0.06), Inches(1.9), fill=NAVY)
    add_text(s, MX + Inches(0.4), box_top + Inches(0.3),
             SLIDE_W - 2 * MX - Inches(0.8), Inches(1.4),
             "Bus bunching causes multiple buses to arrive together, leading to "
             "uneven service and long waiting times for passengers. Manual "
             "tracking makes it hard to monitor buses in real time and respond "
             "quickly. Current systems also lack proper rerouting to handle "
             "traffic or disruptions.",
             16, color=BODY, line_spacing=1.3)


# ---------------------------------------------------------------------------
# SLIDE 3 — Objectives  (EV-deck "01 / green underline" card pattern)
# ---------------------------------------------------------------------------
def numbered_card(s, left, top, w, h, num, title, desc, accent=GREEN):
    add_text(s, left, top, w, Inches(0.3), num, 12, color=GRAYLT,
             name=FONT_BODY)
    hline(s, left, top + Inches(0.34), w, color=accent, weight=Pt(2))
    add_text(s, left, top + Inches(0.5), w, Inches(0.4), title, 16,
             bold=True, color=INK, name=FONT_HEAD)
    add_text(s, left, top + Inches(0.95), w, h - Inches(0.95), desc, 12.5,
             color=BODY, line_spacing=1.2)


def slide_objectives():
    s = new_slide()
    add_title(s, "Objectives")
    add_text(s, MX, Inches(1.35), Inches(9), Inches(0.4),
             "Four core goals guide this system's design", 13, color=BODY)
    items = [
        ("01", "Live Map Tracking",
         "Render bus positions in real time on a Leaflet/OpenStreetMap "
         "interface."),
        ("02", "Bunching Detection",
         "Identify proximity violations and trigger alerts to prevent "
         "clustering."),
        ("03", "Multi-Route Diversion",
         "Present multiple alternative paths to the same destination "
         "during disruptions."),
        ("04", "Admin Dashboard",
         "Enable live operator controls — divert, stop, resume, and "
         "route selection."),
    ]
    col_w = Inches(5.55)
    row_h = Inches(1.7)
    gap_x = Inches(0.6)
    top0 = Inches(2.15)
    for i, (num, title, desc) in enumerate(items):
        col = i % 2
        row = i // 2
        left = MX + col * (col_w + gap_x)
        top = top0 + row * (row_h + Inches(0.3))
        numbered_card(s, left, top, col_w, row_h, num, title, desc)


# ---------------------------------------------------------------------------
# SLIDE 4 — Existing System vs. Proposed System
# ---------------------------------------------------------------------------
def slide_existing_vs_proposed():
    s = new_slide()
    add_title(s, "Existing System vs. Proposed System")
    top = Inches(1.9)
    col_w = Inches(5.6)
    h = Inches(4.2)
    # dark box (existing)
    box1 = rounded_rect(s, MX, top, col_w, h, fill=NAVY, radius=0.035)
    add_text(s, MX + Inches(0.35), top + Inches(0.3), col_w - Inches(0.7),
             Inches(0.5), "Existing System", 18, bold=True, color=WHITE,
             name=FONT_HEAD)
    add_bullets(s, MX + Inches(0.35), top + Inches(0.95), col_w - Inches(0.7),
                Inches(3.0),
                ["Passive GPS tracking with no live decision support",
                 "No visual route diversion or alternative path display",
                 "Weak bunching detection and alert mechanisms",
                 "Manual, delayed response to route disruptions"],
                size=13.5, color=WHITE, bullet_color=RGBColor(0x9A, 0x9D, 0xA6),
                gap=12)
    # light column (proposed)
    left2 = MX + col_w + Inches(0.5)
    add_text(s, left2, top + Inches(0.3), col_w - Inches(0.4), Inches(0.5),
             "Proposed System", 18, bold=True, color=INK, name=FONT_HEAD)
    hline(s, left2, top + Inches(0.85), Inches(1.4), color=GREEN, weight=Pt(2.5))
    add_bullets(s, left2, top + Inches(1.05), col_w - Inches(0.4), Inches(3.0),
                ["Real-time WebSocket updates via FastAPI backend",
                 "React + Leaflet map with live bus positions and route lines",
                 "Admin-driven diversion with multiple alternative routes",
                 "Automated proximity alerts and live activity logging"],
                size=13.5, color=BODY, bullet_color=GREEN, gap=12)


# ---------------------------------------------------------------------------
# SLIDE 5 — System Modules  (plain icon cards, EV "Modules" pattern)
# ---------------------------------------------------------------------------
def icon_card(s, left, top, w, h, symbol, title, desc):
    d = Inches(0.62)
    icon_badge(s, left + d // 2, top + d // 2, d, symbol, fill=CARD,
               sym_color=NAVY)
    add_text(s, left, top + Inches(0.85), w, Inches(0.4), title, 16,
             bold=True, color=INK, name=FONT_HEAD)
    add_text(s, left, top + Inches(1.28), w, h - Inches(1.28), desc, 12.5,
             color=BODY, line_spacing=1.2)


def slide_modules():
    s = new_slide()
    add_title(s, "System Modules")
    items = [
        ("\U0001F5A5", "Backend Simulation",
         "Generates live bus location and speed data for the "
         "Tumkur–Bangalore route."),
        ("\U0001F5B5", "Frontend Dashboard",
         "React UI with map, bus list, alerts, and admin control panel."),
        ("⇄", "WebSocket Communication",
         "Persistent low-latency channel for pushing live updates to the "
         "frontend."),
        ("⚑", "Route Generation & Diversion",
         "Computes and renders alternative paths when a bus is diverted "
         "by admin."),
    ]
    col_w = Inches(5.55)
    row_h = Inches(1.85)
    gap_x = Inches(0.6)
    top0 = Inches(1.85)
    for i, (sym, title, desc) in enumerate(items):
        col = i % 2
        row = i // 2
        left = MX + col * (col_w + gap_x)
        top = top0 + row * (row_h + Inches(0.35))
        icon_card(s, left, top, col_w, row_h, sym, title, desc)


# ---------------------------------------------------------------------------
# SLIDE 6 — System Architecture & Workflow (NEW; EV-deck timeline pattern)
# ---------------------------------------------------------------------------
def slide_architecture():
    s = new_slide()
    add_title(s, "System Architecture and Workflow")
    steps = [
        ("Backend\nSimulation", "Generates bus GPS\nand speed data"),
        ("WebSocket\nStream", "Pushes live ticks\nto the frontend"),
        ("Frontend\nDashboard", "Renders map, bus\nlist, and alerts"),
        ("Bunching\nDetection", "Flags proximity\nviolations"),
        ("Admin\nDecision", "Divert, stop,\nresume, select route"),
        ("Live Map\nUpdate", "Diversion + status\nreflected instantly"),
    ]
    n = len(steps)
    line_y = Inches(3.75)
    left0 = Inches(1.1)
    right0 = SLIDE_W - Inches(1.1)
    hline(s, left0, line_y, right0 - left0, color=BORDER, weight=Pt(1.25))
    step_w = (right0 - left0) / n
    box_d = Inches(0.55)
    for i, (title, desc) in enumerate(steps):
        cx = left0 + step_w * i + step_w / 2
        oval(s, cx - box_d // 2, line_y - box_d // 2, box_d, box_d,
             fill=CARD, line=GREEN, line_w=Pt(1.5))
        above = (i % 2 == 0)
        tw = step_w - Inches(0.15)
        tl = cx - tw / 2
        if above:
            add_text(s, tl, line_y - Inches(1.55), tw, Inches(0.6), title,
                     13, bold=True, color=INK, align=PP_ALIGN.CENTER,
                     name=FONT_HEAD, line_spacing=1.05)
            add_text(s, tl, line_y - Inches(0.95), tw, Inches(0.75), desc,
                     10.5, color=BODY, align=PP_ALIGN.CENTER,
                     line_spacing=1.15)
        else:
            add_text(s, tl, line_y + Inches(0.45), tw, Inches(0.6), title,
                     13, bold=True, color=INK, align=PP_ALIGN.CENTER,
                     name=FONT_HEAD, line_spacing=1.05)
            add_text(s, tl, line_y + Inches(1.05), tw, Inches(0.75), desc,
                     10.5, color=BODY, align=PP_ALIGN.CENTER,
                     line_spacing=1.15)


# ---------------------------------------------------------------------------
# SLIDE 7 — Input & Output
# ---------------------------------------------------------------------------
def slide_io():
    s = new_slide()
    add_title(s, "Input and Output")
    top = Inches(1.9)
    col_w = Inches(5.6)
    add_text(s, MX, top, col_w, Inches(0.4), "System Inputs", 17, bold=True,
             color=INK, name=FONT_HEAD)
    hline(s, MX, top + Inches(0.5), Inches(1.2), color=GREEN, weight=Pt(2.25))
    add_bullets(s, MX, top + Inches(0.75), col_w, Inches(3.2),
                ["Bus GPS coordinates and speed from backend simulation",
                 "Route geometry: waypoints, stops, and destination points",
                 "Admin commands: divert, stop, resume, select route",
                 "Live simulation ticks streamed over WebSocket"],
                size=13.5, color=BODY, bullet_color=GREEN, gap=12)

    left2 = MX + col_w + Inches(0.5)
    add_text(s, left2, top, col_w, Inches(0.4), "System Outputs", 17,
             bold=True, color=INK, name=FONT_HEAD)
    hline(s, left2, top + Inches(0.5), Inches(1.2), color=ORANGE, weight=Pt(2.25))
    add_bullets(s, left2, top + Inches(0.75), col_w, Inches(3.2),
                ["Live bus markers and route polylines on interactive map",
                 "Alternative diversion paths highlighted in distinct colors",
                 "Proximity alerts and real-time bus status indicators",
                 "Activity log and dashboard state updates"],
                size=13.5, color=BODY, bullet_color=ORANGE, gap=12)


# ---------------------------------------------------------------------------
# SLIDE 8 — Mathematical Formulae (mint boxes, EV LP/formula pattern)
# ---------------------------------------------------------------------------
def formula_box(s, left, top, w, h, title, lines):
    rounded_rect(s, left, top, w, h, fill=MINT, radius=0.035)
    add_text(s, left + Inches(0.35), top + Inches(0.28), w - Inches(0.7),
             Inches(0.4), title, 15, bold=True, color=INK, name=FONT_HEAD)
    tb = s.shapes.add_textbox(left + Inches(0.35), top + Inches(0.85),
                               w - Inches(0.7), h - Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(11)
        p.line_spacing = 1.15
        if " = " in line:
            lead, rest = line.split(" = ", 1)
            r1 = p.add_run()
            _set_run(r1, lead, 13, bold=True, color=INK)
            r2 = p.add_run()
            _set_run(r2, " = " + rest, 13, color=INK)
        else:
            r1 = p.add_run()
            _set_run(r1, line, 13, italic=True, color=BODY)


def slide_formulae():
    s = new_slide()
    add_title(s, "Mathematical Formulae — Anti-Bunching Logic")
    add_text(s, MX, Inches(1.3), SLIDE_W - 2 * MX, Inches(0.4),
             "Understanding the core formulas and their application is "
             "crucial for effective bus bunching prevention.", 13,
             color=BODY)
    top = Inches(1.9)
    col_w = Inches(5.6)
    h = Inches(4.3)
    formula_box(s, MX, top, col_w, h, "Core Formulas", [
        "Target Headway = Total Route Distance / Number of Buses",
        "Bunching Threshold = Target Headway × 0.30",
        "Recovery Threshold = Target Headway × 0.60",
        "Holding Threshold = Target Headway × 0.40",
    ])
    left2 = MX + col_w + Inches(0.5)
    formula_box(s, left2, top, col_w, h, "Example Calculation", [
        "For a 20 km route with 3 buses:",
        "Target Headway = 20 / 3 = 6.67 km",
        "Bunching Threshold = 6.67 × 0.30 = 2.00 km",
        "Recovery Threshold = 6.67 × 0.60 = 4.00 km",
        "Holding Threshold = 6.67 × 0.40 = 2.67 km",
    ])
    add_text(s, MX, top + h + Inches(0.15), SLIDE_W - 2 * MX, Inches(0.4),
             "These thresholds guide real-time decisions to prevent "
             "bunching along the route.", 11.5, color=GRAYLT, italic=True)


# ---------------------------------------------------------------------------
# SLIDE 9 — Output Dashboard (real extracted screenshots)
# ---------------------------------------------------------------------------
def slide_dashboard():
    s = new_slide()
    add_title(s, "Output Dashboard")
    map_img = os.path.join(EXTRACTED, "dash_0_1169x684.jpeg")
    list_img = os.path.join(EXTRACTED, "dash_1_438x781.jpeg")

    top = Inches(1.75)
    w1 = Inches(5.5)
    h1 = w1 * 684 / 1169
    rounded_rect(s, MX - Inches(0.06), top - Inches(0.06), w1 + Inches(0.12),
                 h1 + Inches(0.12), fill=None, line=BORDER, line_w=Pt(1),
                 radius=0.03)
    s.shapes.add_picture(map_img, MX, top, width=w1, height=h1)
    add_text(s, MX, top + h1 + Inches(0.1), w1, Inches(0.3),
             "Sample output – 01 (live map)", 11.5, color=GRAYLT,
             italic=True)

    w2 = Inches(2.15)
    h2 = w2 * 781 / 438
    left2 = MX + w1 + Inches(0.45)
    rounded_rect(s, left2 - Inches(0.06), top - Inches(0.06), w2 + Inches(0.12),
                 h2 + Inches(0.12), fill=None, line=BORDER, line_w=Pt(1),
                 radius=0.03)
    s.shapes.add_picture(list_img, left2, top, width=w2, height=h2)
    add_text(s, left2, top + h2 + Inches(0.1), w2, Inches(0.3),
             "Sample output – 02 (bus list)", 11.5, color=GRAYLT,
             italic=True)

    left3 = left2 + w2 + Inches(0.5)
    w3 = SLIDE_W - Inches(0.75) - left3
    add_text(s, left3, top, w3, Inches(0.4), "Dashboard Features", 16,
             bold=True, color=INK, name=FONT_HEAD)
    add_bullets(s, left3, top + Inches(0.55), w3, Inches(4.9), [
        "Live map with bus markers, route line, and stop labels",
        "Connection status indicator (connected / disconnected)",
        "AI Control ON/OFF toggle with live θ / speed-gain tuning sliders",
        "Bunching alert banner when buses cluster together",
        "Live decision feed — hold / no-action / speed-ease / speed-boost",
        "Live metric tiles and rolling CV chart with control-toggle markers",
        "Driver console — full-screen hold/go instructions per bus",
        "Passenger view — live stop ETAs and holding/easing status",
    ], size=11, color=BODY, bullet_color=GREEN, gap=7)


# ---------------------------------------------------------------------------
# SLIDE 10 — Future Work (numbered grid, EV "Completed Phases" pattern)
# ---------------------------------------------------------------------------
def slide_future_work():
    s = new_slide()
    add_title(s, "Future Work")
    items = [
        ("01", "ML-Based Bunching Prediction",
         "Train models to forecast bunching events before they occur "
         "using historical and live data."),
        ("02", "Real GPS Hardware Integration",
         "Replace simulation with physical GPS devices installed on "
         "Tumkur–Bangalore buses."),
        ("03", "Passenger-Facing Mobile App",
         "Extend the system with a public app showing live arrivals and "
         "service alerts."),
        ("04", "Multi-Route & City Scaling",
         "Generalize the platform to support multiple routes and expand "
         "to city-wide deployment."),
        ("05", "Alert & Monitoring",
         "Detects bunching events and logs all system activity in real "
         "time."),
    ]
    col_w = Inches(5.55)
    row_h = Inches(1.45)
    gap_x = Inches(0.6)
    top0 = Inches(1.75)
    for i, (num, title, desc) in enumerate(items):
        col = i % 2
        row = i // 2
        left = MX + col * (col_w + gap_x)
        top = top0 + row * (row_h + Inches(0.25))
        numbered_card(s, left, top, col_w, row_h, num, title, desc,
                      accent=GREEN if i % 2 == 0 else ORANGE)


# ---------------------------------------------------------------------------
# SLIDE 11 — Conclusion (NEW; EV-deck single boxed-statement pattern)
# ---------------------------------------------------------------------------
def slide_conclusion():
    s = new_slide()
    add_title(s, "Conclusion")
    box_top = Inches(2.3)
    rounded_rect(s, MX, box_top, SLIDE_W - 2 * MX, Inches(1.7), fill=WHITE,
                 line=BORDER, line_w=Pt(1), radius=0.04)
    add_text(s, MX + Inches(0.4), box_top + Inches(0.35),
             SLIDE_W - 2 * MX - Inches(0.8), Inches(1.1),
             "Developed an integrated real-time tracking and anti-bunching "
             "dashboard for the Tumkur–Bangalore corridor, combining "
             "live WebSocket telemetry, proximity-based bunching detection, "
             "and admin-driven multi-route diversion in a single operator "
             "interface.", 15.5, color=BODY, line_spacing=1.3)


# ---------------------------------------------------------------------------
# SLIDES 12-13 — References (18 items, split across 2 slides)
# ---------------------------------------------------------------------------
REFERENCES = [
    "Eaton, J. et al., Ant Colony Optimization for Simulated Dynamic Multi-Objective Railway Junction Rescheduling, 2017.",
    "Sun, L. et al., Hybrid Cooperative Co-evolution Algorithm for Uncertain Vehicle Scheduling, 2018.",
    "Deng, W. et al., An Improved Ant Colony Optimization Algorithm Based on Hybrid Strategies for Scheduling Problem, 2019.",
    "Zhang, Q. et al., An Improved Multi-Objective Quantum-Behaved Particle Swarm Optimization for Railway Freight Transportation Routing Design, 2019.",
    "Wu, L. et al., Brainstorming-Based Ant Colony Optimization for Vehicle Routing With Soft Time Windows, 2019.",
    "Liu, W., Route Optimization for Last-Mile Distribution of Rural E-Commerce Logistics Based on ACO, 2020.",
    "Song, Q. et al., Dynamic Path Planning for Unmanned Vehicles Based on Fuzzy Logic and Improved Ant Colony Optimization, 2020.",
    "Wang, Z. et al., Modeling and Planning Multimodal Transport Paths for Risk and Energy Efficiency, 2020.",
    "Liu, D. et al., Two-Echelon Vehicle-Routing Problem: Optimization of Autonomous Delivery Vehicle-Assisted E-Grocery Distribution, 2020.",
    "Zeng, X. et al., Traveling Salesman Problems With Replenishment Arcs and Improved Ant Colony Algorithms, 2021.",
    "Ling, H. et al., An Adaptive Parameter Controlled Ant Colony Optimization Approach for Peer-to-Peer Vehicle and Cargo Matching, 2021.",
    "Lodewijks, G. et al., Reducing CO₂ Emissions of an Airport Baggage Handling Transport System Using a Particle Swarm Optimization Algorithm, 2021.",
    "Li, J. et al., A Multiobjective Ant Colony System Considering Travel and Pollution Costs for Vehicle Routing in Epidemics, 2022.",
    "Govindaraju, S. et al., Intelligent Transportation System's Machine Learning-Based Traffic Prediction, 2024.",
    "Afolayan, B. I. et al., Emerging Trends in Machine Learning Assisted Optimization Techniques Across Intelligent Transportation Systems, 2024.",
    "Xu, P. et al., Ship Formation and Route Optimization Design Based on Improved PSO and D-P Algorithm, 2024.",
    "Hassan, M. et al., Application of Machine Learning in Intelligent Transport Systems: A Comprehensive Review and Bibliometric Analysis, 2025.",
    "Ngo, T. & Van, T., Damage Detection in Bridges by Hybrid Particle Swarm Optimization–Simulated Annealing (PSO-SA) Approach, 2026.",
]


def slide_references(items, title="References"):
    s = new_slide()
    add_title(s, title)
    tb = s.shapes.add_textbox(MX, Inches(1.5), SLIDE_W - 2 * MX, Inches(5.6))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    for i, ref in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        p.line_spacing = 1.15
        r = p.add_run()
        _set_run(r, f"{ref[0]}.  {ref[1]}", 11.5, color=BODY)


def slide_references_split():
    half = (len(REFERENCES) + 1) // 2
    numbered = list(enumerate(REFERENCES, start=1))
    slide_references(numbered[:half], "References")
    slide_references(numbered[half:], "References (continued)")


# ---------------------------------------------------------------------------
def main():
    slide_title()
    slide_problem()
    slide_objectives()
    slide_existing_vs_proposed()
    slide_modules()
    slide_architecture()
    slide_io()
    slide_formulae()
    slide_dashboard()
    slide_future_work()
    slide_conclusion()
    slide_references_split()
    prs.save(OUT)
    print("Saved:", OUT, "| slides:", len(prs.slides))


if __name__ == "__main__":
    main()
